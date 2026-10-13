#!/usr/bin/env python3
"""
scripts/export_docx.py

Converts docs/report.md into a beautifully formatted Microsoft Word (.docx) document
using python-docx.
"""

import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = Path(__file__).resolve().parent.parent
REPORT_MD = BASE_DIR / "docs" / "report.md"
REPORT_DOCX = BASE_DIR / "docs" / "report.docx"

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)

def markdown_to_docx():
    print(f"[INFO] Reading {REPORT_MD}...")
    with open(REPORT_MD, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = Document()

    # Page setup - Margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    lines = md_text.split("\n")
    i = 0
    in_table = False
    table_lines = []

    def flush_table(tbl_lines):
        if not tbl_lines:
            return
        parsed_rows = []
        for l in tbl_lines:
            if re.match(r"^\s*\|?\s*:?-+:?\s*\|", l):
                continue  # separator row
            cells = [c.strip() for c in l.strip().strip("|").split("|")]
            if cells and any(c for c in cells):
                parsed_rows.append(cells)

        if not parsed_rows:
            return

        num_cols = max(len(r) for r in parsed_rows)
        # Pad shorter rows
        for r in parsed_rows:
            while len(r) < num_cols:
                r.append("")

        tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        for row_idx, r in enumerate(parsed_rows):
            is_header = (row_idx == 0)
            for col_idx, cell_value in enumerate(r):
                cell = tbl.cell(row_idx, col_idx)
                # Clean markdown bold/italic
                clean_txt = re.sub(r"\*\*(.*?)\*\*", r"\1", cell_value)
                clean_txt = re.sub(r"\*(.*?)\*", r"\1", clean_txt)
                cell.text = clean_txt
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                if is_header:
                    set_cell_background(cell, "1E3A8A")
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")
        doc.add_paragraph()  # Spacing after table

    while i < len(lines):
        line = lines[i].strip()

        # Handle Tables
        if "|" in line and (line.startswith("|") or line.endswith("|")):
            table_lines.append(line)
            in_table = True
            i += 1
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        if not line:
            i += 1
            continue

        # Handle Headings
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:].strip())
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        elif line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[3:].strip())
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        elif line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[4:].strip())
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        elif line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[5:].strip())
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        elif line.startswith("> "):
            # Callout / Alert
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            txt = line[2:].replace("[!NOTE]", "").replace("[!TIP]", "").replace("[!IMPORTANT]", "").strip()
            run = p.add_run("💬 " + txt)
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            txt = line[2:].strip()
            # Parse bold inside bullet
            parts = re.split(r"(\*\*.*?\*\*)", txt)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        elif re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            txt = re.sub(r"^\d+\.\s+", "", line).strip()
            parts = re.split(r"(\*\*.*?\*\*)", txt)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)
        elif line.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.2
            parts = re.split(r"(\*\*.*?\*\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    p.add_run(part)

        i += 1

    if in_table and table_lines:
        flush_table(table_lines)

    doc.save(REPORT_DOCX)
    print(f"[SUCCESS] Exported Word report to: {REPORT_DOCX}")

if __name__ == "__main__":
    markdown_to_docx()
